import csv
import hmac
import io
import json
import os
import re
import sqlite3
from datetime import datetime
from functools import wraps

from flask import Flask, Response, jsonify, render_template, request, send_file

app = Flask(__name__)
# On Render, set DATABASE_PATH to a file on a persistent disk (e.g. /var/data/audit.db)
# so the data survives restarts and redeploys.
DB = os.environ.get("DATABASE_PATH") or os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "audit.db"
)

# Set ADMIN_PASSWORD in the environment to protect the admin window.
# If it is not set, the admin window is open (fine for local testing only).
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")

BRANCHES = [
    "MANJERI", "KASARGOD", "KANNUR", "KUTTIYADI", "KOZHIKODE", "TIRUR",
    "PALAKKAD", "THRISSUR", "ERNAKULAM", "ALAPPUZHA", "KOTTAYAM", "KOLLAM",
    "TRIVANDRUM", "MARTHANDAM", "NAGPUR", "HYDERABAD", "BANGALORE",
]

# General classroom assets: simple totals.
ASSET = [
    ("student_table", "TOTAL STUDENT TABLE"),
    ("student_chair", "TOTAL STUDENT CHAIR"),
    ("white_board", "TOTAL WHITE BOARD"),
    ("projector", "TOTAL PROJECTOR"),
]

# The three numbered asset tables (SL NO + NAME). Keys are table + SL NO,
# e.g. smt_1, cmt_3, software_6, so the same name can appear in several tables.
TABLES = [
    ("SMT", [
        "LAPTOP", "BLOWER", "IRON", "DC POWER SUPPLY", "MOBILE PHONE DRYER",
        "MINI POLISHING PEN", "MULTI CHARGER", "VIBRATOR", "SEPARATOR MACHINE",
        "LAMINATION MACHINE", "IC STAND", "SHORT KILLER", "MICROSCOPE",
        "MICROSCOPE CAMERA", "BATTERY BOOSTER", "BOOT CABLE", "UV LIGHT",
        "OSCILLOSCOPE", "ROSING PEN", "IPHONE BOARD SEPARATOR",
    ]),
    ("CMT", [
        "LAPTOP", "BGA MACHINE", "BLOWER", "IRON", "SVOD KBC PROGRAMMER",
        "RT809F BIOS PROGRAMMER", "MACBOOK", "IC STAND", "IC STECIL",
        "DC POWER SUPPLY", "C TESTER", "DC CONNECTOR", "PENDRIVE",
        "WORKING MOTHER BOARD", "IDENTIFICATION ITEM", "SCREW DRIVER",
        "MAC SCREW DRIVER SET",
    ]),
    ("SOFTWARE", [
        "UMT", "F64", "UFI", "V1S PRO", "LAPTOP", "SOFTWARE PHONE",
    ]),
]

# key, label (label is what the admin matrix / Excel / CSV show),
# SHORT = the plain name shown inside each card on the branch form.
EQUIPMENT = list(ASSET)
SHORT = dict(ASSET)
GROUPS = [("ASSET", [k for k, _ in ASSET], True)]
for _table, _names in TABLES:
    _keys = []
    for _i, _name in enumerate(_names, 1):
        _key = f"{_table.lower()}_{_i}"
        EQUIPMENT.append((_key, f"{_table} - {_name}"))
        SHORT[_key] = _name
        _keys.append(_key)
    GROUPS.append((_table, _keys, True))
KEYS = [k for k, _ in EQUIPMENT]
GROUP_OF = {k: n for n, ks, _ in GROUPS for k in ks}


def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                branch TEXT NOT NULL,
                classroom_count INTEGER NOT NULL,
                hof_name TEXT NOT NULL,
                hof_contact TEXT NOT NULL,
                counts TEXT NOT NULL,
                lab_assets TEXT,
                suggestions TEXT
            )"""
        )


init_db()


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if ADMIN_PASSWORD:
            auth = request.authorization
            if not auth or not hmac.compare_digest(auth.password or "", ADMIN_PASSWORD):
                return Response(
                    "Admin login required", 401,
                    {"WWW-Authenticate": 'Basic realm="Asset Audit Admin"'},
                )
        return f(*args, **kwargs)
    return wrapper


# ------------------------------------------------------------------ pages
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/admin")
@admin_required
def admin():
    return render_template("admin.html")


# ------------------------------------------------------------------ form API
@app.route("/api/config")
def config():
    return jsonify(
        branches=BRANCHES,
        equipment=[{"key": k, "label": l, "short": SHORT[k]} for k, l in EQUIPMENT],
        groups=[{"name": n, "keys": ks, "numbered": nb} for n, ks, nb in GROUPS],
    )


def validate_submission(d):
    errors = {}

    branch_value = d.get("branch")
    branch = branch_value.strip().upper() if isinstance(branch_value, str) else ""
    if branch not in BRANCHES:
        errors["branch"] = "Choose a branch"

    def to_int(v):
        try:
            n = int(str(v).strip())
            return n if n >= 0 else None
        except (ValueError, TypeError):
            return None

    classroom = to_int(d.get("classroom_count"))
    if classroom is None:
        errors["classroom_count"] = "Enter a valid number"

    hof_value = d.get("hof_name")
    hof = hof_value.strip() if isinstance(hof_value, str) else ""
    if not hof:
        errors["hof_name"] = "Required"

    contact_value = d.get("hof_contact")
    contact = contact_value.strip() if isinstance(contact_value, str) else ""
    if not re.fullmatch(r"\+?[0-9 \-]{8,15}", contact):
        errors["hof_contact"] = "Enter a valid phone number"

    raw_counts = d.get("counts")
    raw_counts = raw_counts if isinstance(raw_counts, dict) else {}
    counts = {}
    for k in KEYS:
        n = to_int(raw_counts.get(k))
        if n is None:
            errors[k] = "Enter a valid number (0 if none)"
        else:
            counts[k] = n

    lab_assets = d.get("lab_assets")
    suggestions = d.get("suggestions")
    values = {
        "branch": branch,
        "classroom_count": classroom,
        "hof_name": hof,
        "hof_contact": contact,
        "counts": counts,
        "lab_assets": lab_assets.strip() if isinstance(lab_assets, str) else "",
        "suggestions": suggestions.strip() if isinstance(suggestions, str) else "",
    }
    return values, errors


@app.route("/api/submit", methods=["POST"])
def submit():
    d = request.get_json(silent=True)
    values, errors = validate_submission(d if isinstance(d, dict) else {})
    if errors:
        return jsonify(ok=False, errors=errors), 400

    with db() as con:
        con.execute(
            "INSERT INTO submissions (created_at, branch, classroom_count, hof_name,"
            " hof_contact, counts, lab_assets, suggestions) VALUES (?,?,?,?,?,?,?,?)",
            (
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                values["branch"], values["classroom_count"], values["hof_name"],
                values["hof_contact"], json.dumps(values["counts"]),
                values["lab_assets"], values["suggestions"],
            ),
        )
    return jsonify(ok=True)


# ------------------------------------------------------------------ admin API
def latest_per_branch():
    """Most recent submission of each branch."""
    with db() as con:
        rows = con.execute(
            "SELECT * FROM submissions WHERE id IN "
            "(SELECT MAX(id) FROM submissions GROUP BY branch)"
        ).fetchall()
    return {r["branch"]: r for r in rows}


def build_matrix(latest):
    matrix = {}
    for k in KEYS:
        matrix[k] = {
            b: (json.loads(latest[b]["counts"]).get(k) if b in latest else None)
            for b in BRANCHES
        }
    return matrix


@app.route("/api/report")
@admin_required
def report():
    latest = latest_per_branch()
    details = {
        b: {
            "id": r["id"],
            "classroom_count": r["classroom_count"],
            "hof_name": r["hof_name"],
            "hof_contact": r["hof_contact"],
            "counts": json.loads(r["counts"]),
            "lab_assets": r["lab_assets"],
            "suggestions": r["suggestions"],
            "submitted": r["created_at"],
        }
        for b, r in latest.items()
    }
    return jsonify(
        branches=BRANCHES,
        equipment=[{"key": k, "label": l, "group": GROUP_OF[k]} for k, l in EQUIPMENT],
        matrix=build_matrix(latest),
        details=details,
        submitted_count=len(latest),
    )


@app.route("/api/submissions/<int:submission_id>", methods=["PUT"])
@admin_required
def update_submission(submission_id):
    d = request.get_json(silent=True)
    if not isinstance(d, dict):
        return jsonify(ok=False, message="Invalid report data."), 400

    values, errors = validate_submission(d)
    if errors:
        return jsonify(ok=False, errors=errors), 400

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db() as con:
        con.execute("BEGIN IMMEDIATE")
        existing = con.execute(
            "SELECT id, branch FROM submissions WHERE id = ?", (submission_id,)
        ).fetchone()
        if existing is None:
            return jsonify(ok=False, message="This report no longer exists."), 404

        latest_source = con.execute(
            "SELECT MAX(id) FROM submissions WHERE branch = ?", (existing["branch"],)
        ).fetchone()[0]
        if latest_source != submission_id:
            return jsonify(
                ok=False,
                message="A newer report was submitted for this branch. Refresh and edit the latest report.",
            ), 409

        latest_target = con.execute(
            "SELECT MAX(id) FROM submissions WHERE branch = ?", (values["branch"],)
        ).fetchone()[0]
        if values["branch"] != existing["branch"] and latest_target is not None and latest_target > submission_id:
            return jsonify(
                ok=False,
                message="A newer report already exists for the selected branch. Choose another branch or refresh.",
            ), 409

        con.execute(
            "UPDATE submissions SET created_at = ?, branch = ?, classroom_count = ?,"
            " hof_name = ?, hof_contact = ?, counts = ?, lab_assets = ?, suggestions = ?"
            " WHERE id = ?",
            (
                now, values["branch"], values["classroom_count"], values["hof_name"],
                values["hof_contact"], json.dumps(values["counts"]), values["lab_assets"],
                values["suggestions"], submission_id,
            ),
        )
    return jsonify(ok=True)


@app.route("/api/export.csv")
@admin_required
def export_csv():
    latest = latest_per_branch()
    matrix = build_matrix(latest)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Equipments"] + BRANCHES)
    for k, label in EQUIPMENT:
        w.writerow([label] + ["" if matrix[k][b] is None else matrix[k][b] for b in BRANCHES])
    w.writerow(["TOTAL"] + [sum(matrix[k][b] or 0 for k in KEYS) for b in BRANCHES])
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=asset_audit_report.csv"},
    )


@app.route("/api/export.xlsx")
@admin_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    latest = latest_per_branch()
    matrix = build_matrix(latest)

    wb = Workbook()
    ws = wb.active
    ws.title = "Asset Audit"
    side = Side(style="thin", color="999999")
    border = Border(left=side, right=side, top=side, bottom=side)
    grey = PatternFill("solid", fgColor="EEEEEE")

    ws.cell(1, 1, "Equipments").font = Font(bold=True, size=14)
    for c, b in enumerate(BRANCHES, 2):
        cell = ws.cell(1, c, b)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1, len(BRANCHES) + 2):
        ws.cell(1, c).fill = grey

    for r, (k, label) in enumerate(EQUIPMENT, 2):
        ws.cell(r, 1, label)
        for c, b in enumerate(BRANCHES, 2):
            v = matrix[k][b]
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(horizontal="center")

    total_row = len(EQUIPMENT) + 2
    ws.cell(total_row, 1, "TOTAL").font = Font(bold=True, size=13)
    for c in range(2, len(BRANCHES) + 2):
        L = get_column_letter(c)
        cell = ws.cell(total_row, c, f"=SUM({L}2:{L}{total_row - 1})")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for c in range(1, len(BRANCHES) + 2):
        ws.cell(total_row, c).fill = grey

    for r in range(1, total_row + 1):
        for c in range(1, len(BRANCHES) + 2):
            ws.cell(r, c).border = border
    ws.column_dimensions["A"].width = 26
    for c in range(2, len(BRANCHES) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "B2"

    # second sheet: branch details
    ds = wb.create_sheet("Branch details")
    ds.append(["Branch", "Class rooms", "HOF name", "HOF contact", "Lab assets", "Suggestions", "Submitted"])
    for c in range(1, 8):
        ds.cell(1, c).font = Font(bold=True)
        ds.cell(1, c).fill = grey
    for b in BRANCHES:
        r = latest.get(b)
        if r:
            ds.append([b, r["classroom_count"], r["hof_name"], r["hof_contact"],
                       r["lab_assets"], r["suggestions"], r["created_at"]])
    for col, w in zip("ABCDEFG", (14, 12, 22, 16, 40, 40, 20)):
        ds.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="asset_audit_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)